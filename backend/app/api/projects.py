from datetime import date

from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload
from openpyxl import load_workbook

from app.core.database import get_db
from app.core.security import verify_password
from app.models import AuditLog, Blocker, Phase, Project, ProjectLevelChange, ProjectMember, User
from app.rbac.deps import (
    require_password_updated,
    get_user_permissions,
    require_permission,
    can_edit_project,
)
from app.schemas import (
    OverviewStats,
    KeyTaskBrief,
    PhaseCreate,
    PhaseOut,
    PhaseUpdate,
    ProjectCreate,
    ProjectDeleteIn,
    ProjectDetail,
    ProjectLevelUpdate,
    ProjectListItem,
    ProjectListOut,
    ProjectMemberOut,
    ProjectMembersUpdate,
    ProjectUpdate,
    RoleCandidateOut,
    UserBrief,
)
from app.services.project_service import (
    cleanup_project_attachment_files,
    delete_project_cascade,
    get_blocker_attribution,
    get_overview_stats,
    get_project_detail,
    list_key_tasks,
    list_projects,
    list_projects_page,
    maybe_replace_phases_for_type_change,
    seed_standard_phases,
    sync_project_phase_progress,
)
from app.services.project_phases import PHASES_BY_TYPE, STANDARD_PHASES, phases_for
from app.services.project_roles import (
    default_roles_for_creator,
    is_project_manager_name,
    is_project_support_user,
)
from app.services.task_notify import notify_project_role
from app.services.user_sync import is_test_account

router = APIRouter(prefix="/projects", tags=["项目"])


def _phase_out(phase: Phase) -> PhaseOut:
    today = date.today()
    overdue = bool(
        phase.planned_end_date
        and phase.planned_end_date < today
        and phase.status != "completed"
    )
    return PhaseOut(
        id=phase.id,
        name=phase.name,
        planned_end_date=phase.planned_end_date,
        status=phase.status,
        note=phase.note,
        sort_order=phase.sort_order,
        is_overdue=overdue,
    )


@router.get("/overview/stats", response_model=OverviewStats)
async def overview_stats(
    user: User = Depends(require_password_updated),
    db: AsyncSession = Depends(get_db),
) -> OverviewStats:
    return await get_overview_stats(db, user, get_user_permissions(user))


@router.get("/overview/key-tasks", response_model=list[KeyTaskBrief])
async def overview_key_tasks(
    user: User = Depends(require_password_updated),
    db: AsyncSession = Depends(get_db),
) -> list[KeyTaskBrief]:
    """总览：未完成的重点任务列表。"""
    projects = await list_projects(db, user, get_user_permissions(user))
    return await list_key_tasks(db, [p.id for p in projects])


@router.get("/overview/blocker-attribution")
async def overview_blocker_attribution(
    user: User = Depends(require_password_updated),
    db: AsyncSession = Depends(get_db),
) -> dict:
    """总览：未关闭卡点按责任方类型归因汇总。"""
    return await get_blocker_attribution(db, user, get_user_permissions(user))


@router.get("", response_model=ProjectListOut)
async def get_projects(
    level: str | None = None,
    project_type: str | None = None,
    status: str | None = None,
    owner_id: int | None = None,
    search: str | None = Query(None),
    page: int = Query(default=1, ge=1, description="页码，从 1 开始"),
    page_size: int = Query(default=10, ge=1, le=100, description="每页条数，默认 10"),
    user: User = Depends(require_password_updated),
    db: AsyncSession = Depends(get_db),
) -> ProjectListOut:
    return await list_projects_page(
        db,
        user,
        get_user_permissions(user),
        level=level,
        project_type=project_type,
        status=status,
        owner_id=owner_id,
        search=search,
        page=page,
        page_size=page_size,
    )


@router.post("", response_model=ProjectDetail)
async def create_project(
    body: ProjectCreate,
    user: User = Depends(require_permission("project.create")),
    db: AsyncSession = Depends(get_db),
) -> ProjectDetail:
    if body.project_type not in {"software", "integration", "fusion"}:
        raise HTTPException(status_code=400, detail="项目类型无效")
    name = (body.name or "").strip()
    if not name:
        raise HTTPException(status_code=400, detail="请填写项目名称")
    if name == "日常任务":
        raise HTTPException(status_code=400, detail="「日常任务」为系统内置名称，请换用真实立项名")

    lead_id = body.lead_user_id
    support_ids = list(dict.fromkeys(body.support_user_ids or []))
    if not support_ids and body.support_user_id is not None:
        support_ids = [body.support_user_id]

    # 创建人身份默认带上：经理→本人为经理；支撑→本人为支撑
    default_lead, default_supports = default_roles_for_creator(user)
    if lead_id is None and default_lead is not None:
        lead_id = default_lead
    if default_supports:
        support_ids = list(dict.fromkeys([*default_supports, *support_ids]))

    if not support_ids:
        raise HTTPException(status_code=400, detail="请至少选择一名项目支撑")
    if lead_id is not None and lead_id in support_ids:
        raise HTTPException(status_code=400, detail="项目经理和项目支撑不能是同一个人")

    all_ids = set(support_ids)
    if lead_id is not None:
        all_ids.add(lead_id)
    found = await db.execute(select(User).where(User.id.in_(all_ids)))
    by_id = {u.id: u for u in found.scalars().all()}
    if set(by_id) != all_ids:
        raise HTTPException(status_code=400, detail="存在无效用户")

    lead_user = by_id.get(lead_id) if lead_id is not None else None
    if lead_user is not None and not is_project_manager_name(lead_user.name):
        raise HTTPException(status_code=400, detail="项目经理只能是许金斌或周永祺")
    for sid in support_ids:
        if not is_project_support_user(by_id[sid]):
            raise HTTPException(
                status_code=400,
                detail="项目支撑须为团队同事（许金斌、周永祺、孙振振除外）",
            )

    # 归属人=创建人（可以是项目支撑，不必是项目经理）
    phase_names = phases_for(body.project_type)
    project = Project(
        name=name,
        project_type=body.project_type,
        level=body.level,
        owner_id=user.id,
        start_date=body.start_date,
        end_date=body.end_date,
        description=body.description,
        current_phase=phase_names[0] if phase_names else STANDARD_PHASES[0],
    )
    db.add(project)
    await db.flush()
    if lead_user is not None:
        db.add(ProjectMember(project_id=project.id, user_id=lead_user.id, member_role="lead"))
    for sid in support_ids:
        db.add(ProjectMember(project_id=project.id, user_id=sid, member_role="support"))
    # 创建人若不在经理/支撑名单里，仍挂为成员，方便自己看得见
    if user.id not in all_ids:
        db.add(ProjectMember(project_id=project.id, user_id=user.id, member_role="member"))
    await seed_standard_phases(db, project.id, body.project_type)
    db.add(AuditLog(user_id=user.id, action="create", resource="project", resource_id=project.id))
    project_id = project.id
    actor_name = user.name or "同事"
    await db.commit()

    if lead_id is not None:
        lead_user = await db.get(User, lead_id)
        if lead_user is not None:
            await notify_project_role(
                db,
                lead_user,
                role_label="项目经理",
                project_name=name,
                project_id=project_id,
                actor_name=actor_name,
            )
    for sid in support_ids:
        support_user = await db.get(User, sid)
        if support_user is not None:
            await notify_project_role(
                db,
                support_user,
                role_label="项目支撑",
                project_name=name,
                project_id=project_id,
                actor_name=actor_name,
            )

    detail = await get_project_detail(db, project_id)
    assert detail is not None
    return detail


@router.get("/role-candidates", response_model=list[RoleCandidateOut])
async def list_role_candidates(
    user: User = Depends(require_password_updated),
    db: AsyncSession = Depends(get_db),
) -> list[RoleCandidateOut]:
    """给新建项目和成员角色表单用：在职人员的姓名和钉钉岗位。"""
    perms = get_user_permissions(user)
    if "project.members" not in perms and "project.create" not in perms:
        raise HTTPException(status_code=403, detail="权限不足")
    result = await db.execute(
        select(User).where(User.is_active.is_(True)).order_by(User.id)
    )
    people = [u for u in result.scalars().all() if not is_test_account(u.username)]
    return [
        RoleCandidateOut(id=u.id, name=u.name, job_title=u.job_title or "")
        for u in people
    ]


@router.get("/{project_id}", response_model=ProjectDetail)
async def get_project(
    project_id: int,
    user: User = Depends(require_password_updated),
    db: AsyncSession = Depends(get_db),
) -> ProjectDetail:
    detail = await get_project_detail(db, project_id)
    if not detail:
        raise HTTPException(status_code=404, detail="项目不存在")
    return detail


@router.patch("/{project_id}", response_model=ProjectDetail)
async def update_project(
    project_id: int,
    body: ProjectUpdate,
    user: User = Depends(require_permission("project.edit")),
    db: AsyncSession = Depends(get_db),
) -> ProjectDetail:
    result = await db.execute(select(Project).where(Project.id == project_id))
    project = result.scalar_one_or_none()
    if not project:
        raise HTTPException(status_code=404, detail="项目不存在")
    if not can_edit_project(user, owner_id=project.owner_id):
        raise HTTPException(status_code=403, detail="仅项目负责人或管理员可修改该项目")
    allowed_types = {"software", "integration", "fusion"}
    data = body.model_dump(exclude_unset=True)
    if "project_type" in data and data["project_type"] not in allowed_types:
        raise HTTPException(status_code=400, detail="项目类型无效")
    if "name" in data and (data["name"] or "").strip() == "日常任务":
        raise HTTPException(status_code=400, detail="「日常任务」为系统内置名称，请换用真实立项名")
    old_type = project.project_type
    new_current_phase = data.pop("current_phase", None) if "current_phase" in data else None
    type_changing = "project_type" in data and data["project_type"] != old_type
    for field, value in data.items():
        setattr(project, field, value)
    if type_changing:
        await maybe_replace_phases_for_type_change(
            db, project, old_type=old_type, new_type=project.project_type
        )
    if new_current_phase is not None:
        await sync_project_phase_progress(
            db, project, phase_name=new_current_phase, status="in_progress"
        )
    await db.commit()
    detail = await get_project_detail(db, project_id)
    assert detail is not None
    return detail


@router.delete("/{project_id}")
async def delete_project(
    project_id: int,
    body: ProjectDeleteIn,
    user: User = Depends(require_permission("project.delete")),
    db: AsyncSession = Depends(get_db),
) -> dict:
    """管理员删除项目。须校验登录密码。先写审计日志，再级联清理关联。"""
    result = await db.execute(select(Project).where(Project.id == project_id))
    project = result.scalar_one_or_none()
    if not project:
        raise HTTPException(status_code=404, detail="项目不存在")
    if project.name == "日常任务":
        raise HTTPException(status_code=400, detail="系统内置项目「日常任务」不可删除")

    try:
        password_ok = verify_password(body.password, user.hashed_password)
    except ValueError:
        password_ok = False
    if not password_ok:
        raise HTTPException(status_code=400, detail="密码不正确")

    project_name = project.name or ""
    project_type = project.project_type or ""
    owner_id = project.owner_id
    db.add(
        AuditLog(
            user_id=user.id,
            action="delete",
            resource="project",
            resource_id=project_id,
            detail=f"名称={project_name}；类型={project_type}；负责人ID={owner_id}",
        )
    )
    stored_paths = await delete_project_cascade(db, project_id)
    await db.commit()
    cleanup_project_attachment_files(stored_paths)
    return {"ok": True}


@router.patch("/{project_id}/level", response_model=ProjectDetail)
async def update_project_level(
    project_id: int,
    body: ProjectLevelUpdate,
    user: User = Depends(require_permission("project.level.change")),
    db: AsyncSession = Depends(get_db),
) -> ProjectDetail:
    result = await db.execute(select(Project).where(Project.id == project_id))
    project = result.scalar_one_or_none()
    if not project:
        raise HTTPException(status_code=404, detail="项目不存在")
    if project.level != body.level:
        db.add(
            ProjectLevelChange(
                project_id=project.id,
                changed_by_id=user.id,
                old_level=project.level,
                new_level=body.level,
            )
        )
        project.level = body.level
    await db.commit()
    detail = await get_project_detail(db, project_id)
    assert detail is not None
    return detail


@router.get("/{project_id}/phases", response_model=list[PhaseOut])
async def list_phases(
    project_id: int,
    user: User = Depends(require_password_updated),
    db: AsyncSession = Depends(get_db),
) -> list[PhaseOut]:
    result = await db.execute(
        select(Phase).where(Phase.project_id == project_id).order_by(Phase.sort_order)
    )
    return [_phase_out(p) for p in result.scalars().all()]


def _member_out(m: ProjectMember) -> ProjectMemberOut:
    return ProjectMemberOut(
        user=UserBrief(
            id=m.user.id,
            name=m.user.name,
            username=m.user.username,
            avatar_url=m.user.avatar_url,
        ),
        member_role=m.member_role or "member",
    )


@router.get("/{project_id}/members", response_model=list[ProjectMemberOut])
async def list_members(
    project_id: int,
    user: User = Depends(require_password_updated),
    db: AsyncSession = Depends(get_db),
) -> list[ProjectMemberOut]:
    result = await db.execute(
        select(ProjectMember)
        .options(selectinload(ProjectMember.user))
        .where(ProjectMember.project_id == project_id)
    )
    members = result.scalars().all()
    return [_member_out(m) for m in members]


@router.put("/{project_id}/members", response_model=list[ProjectMemberOut])
async def update_members(
    project_id: int,
    body: ProjectMembersUpdate,
    user: User = Depends(require_permission("project.members")),
    db: AsyncSession = Depends(get_db),
) -> list[ProjectMemberOut]:
    result = await db.execute(select(Project).where(Project.id == project_id))
    project = result.scalar_one_or_none()
    if not project:
        raise HTTPException(status_code=404, detail="项目不存在")

    lead_id = body.lead_user_id
    support_ids = list(dict.fromkeys(body.support_user_ids or []))
    if not support_ids and body.support_user_id is not None:
        support_ids = [body.support_user_id]
    assistants = set(body.assistant_user_ids)
    members = set(body.member_user_ids)

    if not support_ids:
        raise HTTPException(status_code=400, detail="请至少选择一名项目支撑")
    support_set = set(support_ids)
    if lead_id is not None and lead_id in support_set:
        raise HTTPException(status_code=400, detail="项目经理和项目支撑不能是同一个人")
    if lead_id is not None and (lead_id in assistants or lead_id in members):
        raise HTTPException(status_code=400, detail="项目经理不能同时设为辅助人或普通成员")
    if support_set & assistants or support_set & members:
        raise HTTPException(status_code=400, detail="项目支撑不能同时设为辅助人或普通成员")
    if assistants & members:
        raise HTTPException(status_code=400, detail="同一人不能既是辅助人又是普通成员")

    all_ids = support_set | assistants | members
    if lead_id is not None:
        all_ids.add(lead_id)
    users_result = await db.execute(select(User).where(User.id.in_(all_ids)))
    found_users = list(users_result.scalars().all())
    by_id = {u.id: u for u in found_users}
    if set(by_id) != all_ids:
        raise HTTPException(status_code=400, detail="存在无效用户")

    if lead_id is not None:
        lead_user = by_id[lead_id]
        if not is_project_manager_name(lead_user.name):
            raise HTTPException(status_code=400, detail="项目经理只能是许金斌或周永祺")
    for sid in support_ids:
        if not is_project_support_user(by_id[sid]):
            raise HTTPException(
                status_code=400,
                detail="项目支撑须为团队同事（许金斌、周永祺、孙振振除外）",
            )

    existing = await db.execute(
        select(ProjectMember).where(ProjectMember.project_id == project_id)
    )
    by_user = {m.user_id: m for m in existing.scalars().all()}

    # 删除不在名单中的成员
    for uid, m in list(by_user.items()):
        if uid not in all_ids:
            await db.delete(m)

    def upsert(uid: int, role: str) -> None:
        m = by_user.get(uid)
        if m:
            m.member_role = role
        else:
            db.add(ProjectMember(project_id=project_id, user_id=uid, member_role=role))

    if lead_id is not None:
        upsert(lead_id, "lead")
        project.owner_id = lead_id
    for sid in support_ids:
        upsert(sid, "support")
    # 清空项目经理时不改 owner：归属人仍是创建人或原归属，界面用成员角色显示「未指定」
    for uid in assistants:
        upsert(uid, "assistant")
    for uid in members:
        upsert(uid, "member")

    db.add(
        AuditLog(
            user_id=user.id,
            action="update_members",
            resource="project",
            resource_id=project_id,
            detail=f"lead={lead_id};supports={support_ids};assistants={sorted(assistants)}",
        )
    )
    await db.commit()
    result = await db.execute(
        select(ProjectMember)
        .options(selectinload(ProjectMember.user))
        .where(ProjectMember.project_id == project_id)
    )
    return [_member_out(m) for m in result.scalars().all()]


@router.post("/{project_id}/phases", response_model=PhaseOut)
async def create_phase(
    project_id: int,
    body: PhaseCreate,
    user: User = Depends(require_permission("phase.edit")),
    db: AsyncSession = Depends(get_db),
) -> PhaseOut:
    phase = Phase(project_id=project_id, **body.model_dump())
    db.add(phase)
    await db.commit()
    await db.refresh(phase)
    return _phase_out(phase)


@router.patch("/phases/{phase_id}", response_model=PhaseOut)
async def update_phase(
    phase_id: int,
    body: PhaseUpdate,
    user: User = Depends(require_permission("phase.edit")),
    db: AsyncSession = Depends(get_db),
) -> PhaseOut:
    result = await db.execute(select(Phase).where(Phase.id == phase_id))
    phase = result.scalar_one_or_none()
    if not phase:
        raise HTTPException(status_code=404, detail="阶段不存在")
    project = await db.get(Project, phase.project_id)
    if not project:
        raise HTTPException(status_code=404, detail="项目不存在")
    data = body.model_dump(exclude_unset=True)
    new_status = data.pop("status", None) if "status" in data else None
    for field, value in data.items():
        setattr(phase, field, value)
    if new_status is not None:
        await db.flush()
        await sync_project_phase_progress(
            db, project, phase_id=phase_id, status=new_status
        )
        await db.commit()
        result = await db.execute(select(Phase).where(Phase.id == phase_id))
        phase = result.scalar_one()
        return _phase_out(phase)
    await db.commit()
    await db.refresh(phase)
    return _phase_out(phase)


@router.post("/import/excel")
async def import_projects_excel(
    file: UploadFile = File(...),
    user: User = Depends(require_permission("project.create")),
    db: AsyncSession = Depends(get_db),
) -> dict:
    wb = load_workbook(file.file)
    ws = wb.active
    created = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        name, project_type, level, owner_username = row[0], row[1] or "software", row[2] or "normal", row[3]
        owner_result = await db.execute(select(User).where(User.username == str(owner_username)))
        owner = owner_result.scalar_one_or_none()
        if not owner:
            continue
        raw_level = str(level).strip().lower()
        level_map = {
            "green": "normal",
            "yellow": "key",
            "red": "key",
            "普通": "normal",
            "重点": "key",
            "日常": "daily",
            "normal": "normal",
            "key": "key",
            "daily": "daily",
        }
        level_val = level_map.get(raw_level, "normal")
        raw_type = str(project_type).strip().lower()
        type_map = {
            "software": "software",
            "软件开发": "software",
            "软件开发项目": "software",
            "integration": "integration",
            "系统集成": "integration",
            "系统集成项目": "integration",
            "fusion": "fusion",
            "集众融合": "fusion",
            "集众融合项目": "fusion",
            # 历史类型映射
            "ops": "integration",
            "engineering": "integration",
            "procurement": "integration",
            "hardware": "integration",
            "operation": "fusion",
            "工程装修": "integration",
            "采购": "integration",
            "硬件对接": "integration",
            "运维": "integration",
            "运营": "fusion",
        }
        type_val = type_map.get(raw_type, "software")
        if "智安校园" in str(name):
            type_val = "fusion"
        project = Project(
            name=str(name),
            project_type=type_val,
            level=level_val,
            owner_id=owner.id,
        )
        db.add(project)
        await db.flush()
        db.add(ProjectMember(project_id=project.id, user_id=owner.id, member_role="lead"))
        created += 1
    await db.commit()
    return {"created": created}


@router.get("/meta/types")
async def project_types() -> dict:
    return {
        "types": [
            {"value": "software", "label": "软件开发项目"},
            {"value": "integration", "label": "系统集成项目"},
            {
                "value": "fusion",
                "label": "集众融合项目",
                "hint": "如智安校园等融合类产品",
            },
        ],
        "phases": [{"value": p, "label": p} for p in STANDARD_PHASES],
        "phases_by_type": {
            key: [{"value": p, "label": p} for p in names]
            for key, names in PHASES_BY_TYPE.items()
        },
        "levels": [
            {"value": "normal", "label": "普通"},
            {"value": "key", "label": "重点"},
            {"value": "daily", "label": "日常"},
        ],
        "task_priorities": [
            {"value": "p0", "label": "P0", "hint": "最高，会提醒责任人"},
            {"value": "p1", "label": "P1"},
            {"value": "p2", "label": "P2"},
            {"value": "p3", "label": "P3", "hint": "默认"},
            {"value": "p4", "label": "P4", "hint": "最低"},
        ],
        "statuses": [
            {"value": "normal", "label": "正常", "emoji": "🟢"},
            {"value": "at_risk", "label": "有风险", "emoji": "🟡"},
            {"value": "blocked", "label": "已阻塞", "emoji": "🔴"},
        ],
    }
