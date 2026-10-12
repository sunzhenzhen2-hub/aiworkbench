"""从方案文件抽取文本（PDF / Word / PPT）；图片交由多模态模型处理。"""

from __future__ import annotations

from pathlib import Path

from app.services.solution_storage import is_image_filename

# 送给模型的文本上限，避免超长上下文
MAX_EXTRACT_CHARS = 12000


def extract_text_from_file(path: Path, original_filename: str) -> str:
    """抽取文档正文；图片返回空字符串（由 vision 路径处理）。"""
    if is_image_filename(original_filename) or is_image_filename(path.name):
        return ""

    ext = path.suffix.lower()
    if ext == ".pdf":
        return _extract_pdf(path)
    if ext == ".docx":
        return _extract_docx(path)
    if ext == ".pptx":
        return _extract_pptx(path)
    return ""


def _truncate(text: str) -> str:
    text = (text or "").strip()
    if len(text) <= MAX_EXTRACT_CHARS:
        return text
    return text[:MAX_EXTRACT_CHARS] + "\n…（后续已截断）"


def _extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    parts: list[str] = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:  # noqa: BLE001
            continue
    return _truncate("\n".join(parts))


def _extract_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    # 表格单元格
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text and cell.text.strip():
                    parts.append(cell.text.strip())
    return _truncate("\n".join(parts))


def _extract_pptx(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(path))
    parts: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text and shape.text.strip():
                parts.append(shape.text.strip())
    return _truncate("\n".join(parts))
