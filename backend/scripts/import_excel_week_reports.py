"""将 docs 下 Excel 周报入库：LLM 结构化为三问卡片后写入 progress_reports。

用法（在 backend 目录）：
  PYTHONPATH=. .venv/bin/python scripts/import_excel_week_reports.py
"""

from __future__ import annotations

import asyncio
import json
import re
import sys
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

import httpx
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import selectinload

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from app.core.database import async_session  # noqa: E402
from app.models import ProgressReport, ProgressReportItem, User  # noqa: E402
from app.services.llm_resolver import resolve_llm  # noqa: E402
from app.services.project_service import backfill_null_report_projects  # noqa: E402
from app.services.report_ai_structure import _item_to_structured  # noqa: E402
from app.services.work_tags import normalize_work_tags  # noqa: E402

PERIOD_START = date(2026, 9, 7)
PERIOD_END = date(2026, 9, 11)
PERIOD_TYPE = "week"

EXCEL_DIR = Path(__file__).resolve().parents[2] / "docs/9月7日-11日周报汇总/9月7日-11日周报汇总"

# 文件名关键词 → 系统姓名（须与 users.name 一致）
FILE_NAME_MAP: list[tuple[str, str]] = [
    ("许金斌", "许金斌"),
    ("徐初", "徐初"),
    ("孙聪", "孙聪"),
    ("钱波", "钱波"),
    ("朱高见", "朱高见"),
    ("金滨", "金滨"),
    ("凌亮", "凌亮"),
    ("周永祺", "周永祺"),
    ("罗琳", "罗琳"),
    ("王捷", "王捷"),
]

SKIP_SHEET_HINTS = ("投标", "客户侧", "项目进展", "副本")


def _cell(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.isoformat()
    return str(v).strip()


def _is_noise_row(texts: list[str]) -> bool:
    joined = " ".join(texts)
    if not joined.strip():
        return True
    # 表头行
    if "工作事项" in joined and "本周进展" in joined:
        return True
    if "工作事项" in joined and "进展说明" in joined:
        return True
    # 纯占位
    meaningful = [t for t in texts if t and t not in {"/", "／", "-", "—", "无", "无进度", "持续中"}]
    return len(meaningful) == 0


def extract_xlsx(path: Path) -> str:
    wb = load_workbook(path, data_only=True, read_only=True)
    chunks: list[str] = []
    for ws in wb.worksheets:
        title = ws.title or ""
        if any(h in title for h in SKIP_SHEET_HINTS):
            continue
        rows_out: list[str] = []
        for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if i > 120:
                break
            cells = [_cell(c) for c in row]
            # 去掉全空尾列
            while cells and not cells[-1]:
                cells.pop()
            if _is_noise_row(cells):
                continue
            # 压缩：只保留有内容的列，带列号便于 LLM 理解
            parts = [f"[{j}]{c}" for j, c in enumerate(cells) if c]
            if parts:
                rows_out.append(f"R{i}: " + " | ".join(parts))
        if rows_out:
            chunks.append(f"## 工作表: {title}\n" + "\n".join(rows_out[:80]))
    wb.close()
    return "\n\n".join(chunks)


def extract_xls(path: Path) -> str:
    import xlrd

    book = xlrd.open_workbook(str(path))
    chunks: list[str] = []
    for si in range(book.nsheets):
        sh = book.sheet_by_index(si)
        title = sh.name or f"sheet{si}"
        if any(h in title for h in SKIP_SHEET_HINTS):
            continue
        rows_out: list[str] = []
        for i in range(min(sh.nrows, 80)):
            cells = [_cell(sh.cell_value(i, j)) for j in range(sh.ncols)]
            while cells and not cells[-1]:
                cells.pop()
            if _is_noise_row(cells):
                continue
            parts = [f"[{j}]{c}" for j, c in enumerate(cells) if c]
            if parts:
                rows_out.append(f"R{i + 1}: " + " | ".join(parts))
        if rows_out:
            chunks.append(f"## 工作表: {title}\n" + "\n".join(rows_out[:60]))
    return "\n\n".join(chunks)


def extract_file(path: Path) -> str:
    if path.suffix.lower() == ".xls":
        return extract_xls(path)
    return extract_xlsx(path)


def _extract_json_object(text: str) -> dict[str, Any]:
    """从模型回复中抠出 JSON 对象（兼容 ```json 包裹与思考标签）。"""
    cleaned = text or ""
    cleaned = re.sub(r"<think>[\s\S]*?</think>", "", cleaned)
    cleaned = cleaned.strip()
    if cleaned.startswith("```"):
        cleaned = re.sub(r"^```(?:json)?\s*", "", cleaned)
        cleaned = re.sub(r"\s*```$", "", cleaned)
    start = cleaned.find("{")
    end = cleaned.rfind("}")
    if start < 0 or end < 0 or end <= start:
        raise ValueError(f"LLM 未返回 JSON: {(text or '')[:300]}")
    return json.loads(cleaned[start : end + 1])


async def llm_structure(raw_text: str, person_name: str, llm) -> dict[str, Any]:
    """一次调用：整份 Excel 原文 → theme + items[]。"""
    system = (
        "你是项目周报结构化助手。把员工 Excel 周报原文整理成系统三问卡片。\n"
        "只输出 JSON（不要 markdown），格式：\n"
        "{"
        '"theme":"本周主题≤40字",'
        '"items":[{'
        '"task_title":"事项名≤120字",'
        '"result_text":"本周做成了什么（可验证）",'
        '"quant_value":"数量/完成度，无则空串",'
        '"stuck_status":"none|blocker|coordination",'
        '"stuck_detail":"卡住原因，无则空",'
        '"stuck_next_action":"下一步动作，无则空",'
        '"stuck_follow_up":"YYYY-MM-DD或null",'
        '"stuck_need_who":"需要谁，无则空",'
        '"next_goal":"下周目标",'
        '"next_due_date":"YYYY-MM-DD或null",'
        '"risk_level":"normal|watch|escalate",'
        '"risk_reason":"风险说明，无则空",'
        '"risk_impact_date":null,'
        '"risk_coordinator":"",'
        '"tags":["售前"]'
        "}]}\n"
        "规则：\n"
        "1) 只保留本周（2026-09-07～09-11）真正有进展或明确卡点的事项；合并重复行。\n"
        "2) 跳过空白、仅写「无/无进度/持续中/」且无实质信息的行。\n"
        "3) 跳过多年投标/客户侧历史汇总表内容。\n"
        "4) result_text 必须具体；若原文含糊，用「原文信息不足：…」标注，不要编造。\n"
        "5) 有卡点写 stuck_status=blocker 或 coordination，并写 need_who / next_action。\n"
        "6) 有下周计划写入 next_goal；没有则写「原文未写下周计划」。\n"
        "7) 每人最多 12 条卡片；优先有实质进展的。\n"
        "8) tags 必填 1～3 个，只能从："
        "售前、售中、售后、运维、项目管理、流程跟进、客户对接、方案汇报；"
        "可参考原文「业务域」列；用中文。\n"
        "9) 用中文。"
    )
    # 控制长度，避免超上下文
    body = raw_text[:18000]
    user_msg = (
        f"员工姓名：{person_name}\n"
        f"周期：{PERIOD_START.isoformat()} ~ {PERIOD_END.isoformat()}\n"
        f"Excel 原文：\n{body}"
    )
    url = llm.api_base.rstrip("/") + "/chat/completions"
    async with httpx.AsyncClient(timeout=120.0) as client:
        resp = await client.post(
            url,
            headers={"Authorization": f"Bearer {llm.api_key}"},
            json={
                "model": llm.model_name,
                "messages": [
                    {"role": "system", "content": system},
                    {"role": "user", "content": user_msg},
                ],
                "temperature": 0.2,
            },
        )
        if resp.status_code >= 400:
            raise ValueError(f"LLM HTTP {resp.status_code}: {resp.text[:500]}")
        payload = resp.json()
        try:
            content = payload["choices"][0]["message"]["content"]
        except (KeyError, IndexError, TypeError) as e:
            raise ValueError(f"LLM 响应异常: {str(payload)[:500]}") from e
        if not content:
            # 部分模型把正文放在 reasoning / content 列表里
            msg = payload["choices"][0]["message"]
            content = msg.get("reasoning_content") or msg.get("reasoning") or ""
            if isinstance(content, list):
                content = "".join(
                    (c.get("text") if isinstance(c, dict) else str(c)) for c in content
                )
            if not content:
                raise ValueError(f"LLM 返回空内容: {str(payload)[:500]}")
    return _extract_json_object(content)

def find_excel_files() -> list[tuple[Path, str]]:
    found: list[tuple[Path, str]] = []
    used: set[str] = set()
    for path in sorted(EXCEL_DIR.iterdir()):
        if path.suffix.lower() not in {".xlsx", ".xls"} or path.name.startswith("~"):
            continue
        matched = None
        for key, name in FILE_NAME_MAP:
            if key in path.name:
                matched = name
                break
        if not matched:
            print(f"  [跳过未识别] {path.name}")
            continue
        if matched in used:
            print(f"  [跳过重复] {path.name} -> {matched}")
            continue
        used.add(matched)
        found.append((path, matched))
    return found


async def upsert_report(
    db,
    user: User,
    theme: str,
    items: list[dict[str, Any]],
) -> ProgressReport:
    q = await db.execute(
        select(ProgressReport)
        .options(selectinload(ProgressReport.items))
        .where(
            ProgressReport.user_id == user.id,
            ProgressReport.period_type == PERIOD_TYPE,
            ProgressReport.period_start == PERIOD_START,
        )
    )
    report = q.scalar_one_or_none()
    if report is None:
        report = ProgressReport(
            user_id=user.id,
            period_type=PERIOD_TYPE,
            period_start=PERIOD_START,
            period_end=PERIOD_END,
            theme=theme[:80],
            status="submitted",
            submitted_at=datetime.now(timezone.utc),
        )
        db.add(report)
        await db.flush()
    else:
        report.theme = theme[:80]
        report.period_end = PERIOD_END
        report.status = "submitted"
        report.submitted_at = datetime.now(timezone.utc)
        report.items.clear()
        await db.flush()

    for i, raw in enumerate(items):
        st = _item_to_structured(raw)
        # 兜底枚举
        if st["stuck_status"] not in {"none", "blocker", "coordination"}:
            st["stuck_status"] = "none"
        if st["risk_level"] not in {"normal", "watch", "escalate"}:
            st["risk_level"] = "normal"
        db.add(
            ProgressReportItem(
                report_id=report.id,
                task_title=st["task_title"] or f"事项{i + 1}",
                result_text=st["result_text"],
                quant_value=st["quant_value"],
                stuck_status=st["stuck_status"],
                stuck_detail=st["stuck_detail"],
                stuck_next_action=st["stuck_next_action"],
                stuck_follow_up=st["stuck_follow_up"],
                stuck_need_who=st["stuck_need_who"],
                next_goal=st["next_goal"],
                next_due_date=st["next_due_date"],
                risk_level=st["risk_level"],
                risk_reason=st["risk_reason"],
                risk_impact_date=st["risk_impact_date"],
                risk_coordinator=st["risk_coordinator"],
                tags=normalize_work_tags(raw.get("tags")),
                sort_order=i,
            )
        )
    await db.flush()
    return report


async def main() -> None:
    if not EXCEL_DIR.is_dir():
        raise SystemExit(f"目录不存在: {EXCEL_DIR}")

    files = find_excel_files()
    print(f"待导入 {len(files)} 份 Excel，周期 {PERIOD_START} ~ {PERIOD_END}")

    async with async_session() as db:
        llm = await resolve_llm(db, "report_structure") or await resolve_llm(db, "report_coach")
        if not llm:
            raise SystemExit("未配置可用 LLM（report_structure / report_coach）")
        print(f"使用模型: {llm.name} / {llm.model_name}")

        users = (await db.execute(select(User))).scalars().all()
        # 先固化成普通 dict，避免 rollback 后 ORM 过期触发 MissingGreenlet
        people: dict[str, dict[str, Any]] = {
            u.name: {"id": u.id, "username": u.username, "name": u.name} for u in users
        }

        summary: list[str] = []
        for path, person_name in files:
            person = people.get(person_name)
            if not person:
                msg = f"FAIL {person_name}: 系统无此用户"
                print(msg)
                summary.append(msg)
                continue
            print(f"\n=== {person_name} ({person['username']}) ← {path.name} ===")
            try:
                raw = extract_file(path)
                if not raw.strip():
                    raise ValueError("未能从 Excel 抽出有效行")
                print(f"  原文约 {len(raw)} 字，调用 LLM…")
                data = None
                last_err: Exception | None = None
                for attempt in range(1, 3):
                    try:
                        data = await llm_structure(raw, person_name, llm)
                        break
                    except Exception as e:
                        last_err = e
                        print(f"  LLM 第{attempt}次失败: {type(e).__name__}: {e}")
                        await asyncio.sleep(2)
                if data is None:
                    raise last_err or RuntimeError("LLM 失败")
                theme = (data.get("theme") or f"{person_name} 9/7-9/11 周报").strip()
                items = data.get("items") or []
                if not isinstance(items, list) or not items:
                    raise ValueError(f"LLM 未返回 items: {str(data)[:300]}")
                # 过滤空卡
                items = [
                    it
                    for it in items
                    if isinstance(it, dict)
                    and (
                        (it.get("task_title") or "").strip()
                        or (it.get("result_text") or "").strip()
                    )
                ][:12]
                user = await db.get(User, person["id"])
                if not user:
                    raise ValueError("用户记录丢失")
                report = await upsert_report(db, user, theme, items)
                await db.commit()
                msg = (
                    f"OK {person_name}: report_id={report.id} "
                    f"items={len(items)} theme={theme}"
                )
                print(f"  {msg}")
                summary.append(msg)
            except Exception as e:
                await db.rollback()
                msg = f"FAIL {person_name}: {type(e).__name__}: {e}"
                print(f"  {msg}")
                summary.append(msg)

        if "俞烨涛" in people:
            summary.append("SKIP 俞烨涛: 文件夹中无对应 Excel")

        bf = await backfill_null_report_projects(db)
        await db.commit()
        print(f"\nbackfill_projects={bf}")

        print("\n==== 汇总 ====")
        for s in summary:
            print(s)

        cnt = await db.execute(
            select(ProgressReport).where(
                ProgressReport.period_type == PERIOD_TYPE,
                ProgressReport.period_start == PERIOD_START,
                ProgressReport.status == "submitted",
            )
        )
        reports = cnt.scalars().all()
        print(f"\n本周已提交周报: {len(reports)} 份")

if __name__ == "__main__":
    asyncio.run(main())
