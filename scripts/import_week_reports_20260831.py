#!/usr/bin/env python3
"""将 docs 下 8/31–9/4 Excel 周报优化后写入 progress_reports。

用法（在 backend 目录）:
  PYTHONPATH=. .venv/bin/python ../scripts/import_week_reports_20260831.py
"""

from __future__ import annotations

import re
import sys
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import selectinload

# 保证可从 scripts/ 直接跑
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

from app.core.database import async_session  # noqa: E402
from app.models import ProgressReport, ProgressReportItem, User  # noqa: E402
from app.services.project_service import (  # noqa: E402
    backfill_null_report_projects,
    ensure_daily_project,
)
from app.services.report_quality import (  # noqa: E402
    _is_vague,
    compose_theme_from_titles,
    validate_report_items,
)

REPORT_DIR = ROOT / "docs" / "8月31-9月4日周报汇总" / "8月31-9月4日周报汇总"
PERIOD_START = date(2026, 8, 31)
PERIOD_END = date(2026, 9, 4)
DEFAULT_NEXT_DUE = date(2026, 9, 11)

# 文件名关键词 → 系统姓名（与 users.name 一致）
FILE_AUTHOR_HINTS: list[tuple[str, str]] = [
    ("许金斌", "许金斌"),
    ("俞烨涛", "俞烨涛"),
    ("孙聪", "孙聪"),
    ("徐初", "徐初"),
    ("朱高见", "朱高见"),
    ("钱波", "钱波"),
    ("金滨", "金滨"),
    ("周永祺", "周永祺"),
    ("凌亮", "凌亮"),
    ("罗琳", "罗琳"),
    ("王捷", "王捷"),
]

EMPTYISH = re.compile(
    r"^[\s/／\-—_]*$|^(无|无进展|暂无进度|本周无进度|本周暂无进度|无进度|/|——|-)$"
)
FOLLOW_ONLY = re.compile(r"^(跟进项目进度|跟进|推进中|处理中|持续中|正常)$")


def _cell(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.isoformat()
    return str(v).strip()


def _is_emptyish(text: str) -> bool:
    t = (text or "").strip()
    if not t:
        return True
    if EMPTYISH.match(t) or FOLLOW_ONLY.match(t):
        return True
    return False


def _guess_author_from_filename(name: str) -> str | None:
    for key, author in FILE_AUTHOR_HINTS:
        if key in name:
            return author
    return None


def _pick_sheet_name(names: list[str]) -> str:
    for n in names:
        if "周进度" in n:
            return n
    return names[0]


def _read_xlsx_rows(path: Path) -> tuple[list[list[Any]], str]:
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        sheet = _pick_sheet_name(wb.sheetnames)
        ws = wb[sheet]
        rows: list[list[Any]] = []
        for row in ws.iter_rows(max_col=6, values_only=True):
            rows.append(list(row))
        return rows, sheet
    finally:
        wb.close()


def _read_xls_rows(path: Path) -> tuple[list[list[Any]], str]:
    import xlrd

    book = xlrd.open_workbook(path)
    # 优先周进度表
    sheet = None
    for s in book.sheet_names():
        if "周进度" in s:
            sheet = book.sheet_by_name(s)
            break
    if sheet is None:
        sheet = book.sheet_by_index(0)
    rows = []
    for r in range(sheet.nrows):
        rows.append([sheet.cell_value(r, c) for c in range(min(6, sheet.ncols))])
    return rows, sheet.name


def _detect_header(rows: list[list[Any]]) -> dict[str, int] | None:
    """返回列索引：domain/project/result/next/issue/author。"""
    for row in rows[:5]:
        cells = [_cell(c) for c in row]
        joined = "|".join(cells)
        if "具体项目" in joined or ("本周工作" in joined and "下周" in joined):
            mapping: dict[str, int] = {}
            for i, c in enumerate(cells):
                if "业务域" in c:
                    mapping["domain"] = i
                elif "具体项目" in c or c == "项目":
                    mapping["project"] = i
                elif "本周工作" in c:
                    mapping["result"] = i
                elif "下周" in c:
                    mapping["next"] = i
                elif "存在问题" in c or "协调" in c:
                    mapping["issue"] = i
                elif "汇报人" in c:
                    mapping["author"] = i
            if "project" in mapping and "result" in mapping:
                mapping.setdefault("domain", 0)
                mapping.setdefault("next", mapping["result"] + 1)
                mapping.setdefault("issue", mapping.get("next", 3) + 1)
                return mapping
    return None


def _row_dict(row: list[Any], cols: dict[str, int]) -> dict[str, str]:
    def get(key: str) -> str:
        i = cols.get(key)
        if i is None or i >= len(row):
            return ""
        return _cell(row[i])

    return {
        "domain": get("domain"),
        "project": get("project"),
        "result": get("result"),
        "next": get("next"),
        "issue": get("issue"),
        "author": get("author"),
    }


def _optimize_item(
    *,
    title: str,
    result: str,
    next_goal: str,
    issue: str,
    domain: str = "",
) -> dict[str, Any] | None:
    title = re.sub(r"^\d+[\.、．]\s*", "", title).strip()[:120]
    result = result.strip()
    next_goal = next_goal.strip()
    issue = issue.strip()

    if _is_emptyish(title) and _is_emptyish(result):
        return None
    if not title:
        title = (domain or "未命名工作")[:120]

    # 占位行
    if title in {"/", "／", "-", "——"} and _is_emptyish(result):
        return None

    no_action = False
    if _is_emptyish(result) or _is_vague(result):
        # 若有下周计划或问题，仍保留卡片但标无实质动作；否则跳过纯空卡
        if _is_emptyish(next_goal) and _is_emptyish(issue):
            return None
        no_action = True
        if _is_emptyish(result):
            result = "本周无实质动作（由原周报导入）"

    if no_action:
        return {
            "project_id": None,
            "task_title": title,
            "no_action": True,
            "result_text": result[:2000],
            "quant_value": "",
            "stuck_status": "none",
            "stuck_detail": "",
            "stuck_next_action": "",
            "stuck_follow_up": None,
            "stuck_need_who": "",
            "linked_blocker_id": None,
            "next_goal": "",
            "next_due_date": None,
            "risk_level": "normal",
            "risk_reason": "",
            "risk_impact_date": None,
            "risk_coordinator": "",
            "prior_done": None,
            "sort_order": 0,
        }

    # 有实质动作
    if _is_emptyish(next_goal) or _is_vague(next_goal):
        next_goal = f"继续推进「{title[:40]}」相关工作"

    stuck_status = "none"
    stuck_detail = ""
    stuck_next_action = ""
    stuck_follow_up = None
    if issue and not _is_emptyish(issue) and issue not in {"无", "/", "——"}:
        stuck_status = "coordination"
        stuck_detail = issue[:500]
        stuck_next_action = (next_goal or "跟进协调")[:500]
        stuck_follow_up = DEFAULT_NEXT_DUE

    # 从结果抽简单量化
    quant = ""
    m = re.search(r"(\d+)\s*路", result)
    if m:
        quant = f"{m.group(1)}路"
    else:
        m = re.search(r"(\d+)\s*台", result)
        if m:
            quant = f"{m.group(1)}台"
        else:
            m = re.search(r"约?\s*(\d+)\s*台", result)
            if m:
                quant = f"{m.group(1)}台"

    risk = "normal"
    risk_reason = ""
    if any(k in f"{result}{issue}" for k in ("滞后", "超期", "离线", "无法确定")):
        risk = "watch"
        risk_reason = (issue or result)[:200]

    return {
        "project_id": None,
        "task_title": title,
        "no_action": False,
        "result_text": result[:4000],
        "quant_value": quant[:64],
        "stuck_status": stuck_status,
        "stuck_detail": stuck_detail,
        "stuck_next_action": stuck_next_action,
        "stuck_follow_up": stuck_follow_up,
        "stuck_need_who": "",
        "linked_blocker_id": None,
        "next_goal": next_goal[:2000],
        "next_due_date": DEFAULT_NEXT_DUE,
        "risk_level": risk,
        "risk_reason": risk_reason,
        "risk_impact_date": None,
        "risk_coordinator": "",
        "prior_done": None,
        "sort_order": 0,
    }


def parse_standard_rows(
    rows: list[list[Any]], cols: dict[str, int]
) -> tuple[list[dict[str, Any]], str | None, int]:
    items: list[dict[str, Any]] = []
    author: str | None = None
    skipped = 0
    last_domain = ""
    last_project = ""

    # skip header row(s)
    start = 0
    for i, row in enumerate(rows[:5]):
        cells = [_cell(c) for c in row]
        if "具体项目" in "|".join(cells) or "本周工作" in "|".join(cells):
            start = i + 1
            break

    for row in rows[start:]:
        d = _row_dict(row, cols)
        if d["author"] and not author:
            author = d["author"]
        if d["domain"]:
            last_domain = d["domain"]
        project = d["project"] or last_project
        if d["project"]:
            last_project = d["project"]

        # 纯业务域标题行
        if d["domain"] and not d["project"] and _is_emptyish(d["result"]) and _is_emptyish(d["next"]):
            skipped += 1
            continue

        item = _optimize_item(
            title=project or d["domain"],
            result=d["result"],
            next_goal=d["next"],
            issue=d["issue"],
            domain=last_domain,
        )
        if not item:
            skipped += 1
            continue
        items.append(item)

    # 合并连续同名标题（结果拼接）
    merged: list[dict[str, Any]] = []
    for it in items:
        if (
            merged
            and merged[-1]["task_title"] == it["task_title"]
            and not merged[-1]["no_action"]
            and not it["no_action"]
        ):
            prev = merged[-1]
            prev["result_text"] = (prev["result_text"] + "；" + it["result_text"])[:4000]
            if it["next_goal"] and it["next_goal"] not in prev["next_goal"]:
                prev["next_goal"] = (prev["next_goal"] + "；" + it["next_goal"])[:2000]
            if it["stuck_status"] != "none" and prev["stuck_status"] == "none":
                prev["stuck_status"] = it["stuck_status"]
                prev["stuck_detail"] = it["stuck_detail"]
                prev["stuck_next_action"] = it["stuck_next_action"]
                prev["stuck_follow_up"] = it["stuck_follow_up"]
            continue
        merged.append(it)

    for i, it in enumerate(merged):
        it["sort_order"] = i
    return merged, author, skipped


def parse_jinbin(path: Path) -> tuple[list[dict[str, Any]], str | None, int]:
    """金滨运维周志：日期 / 本周工作进度 / 下周工作计划 / 备注。取 2026 表中 8.31–9.4 行。"""
    import xlrd

    book = xlrd.open_workbook(path)
    sheet = None
    for name in book.sheet_names():
        if "2026" in name:
            sheet = book.sheet_by_name(name)
            break
    if sheet is None:
        sheet = book.sheet_by_index(-1 if book.nsheets > 1 else 0)

    target_keys = ("8月31日至9月4日", "8月31日-9月4日", "8.31-9.4", "8月31日至9月6日")
    hit = None
    for r in range(sheet.nrows):
        date_cell = str(sheet.cell_value(r, 0)).strip().replace(" ", "")
        if any(k.replace(" ", "") in date_cell for k in target_keys):
            hit = (
                date_cell,
                str(sheet.cell_value(r, 1)).strip() if sheet.ncols > 1 else "",
                str(sheet.cell_value(r, 2)).strip() if sheet.ncols > 2 else "",
                str(sheet.cell_value(r, 3)).strip() if sheet.ncols > 3 else "",
            )
            break
    if not hit or not hit[1]:
        return [], "金滨", 0

    _date, result, next_goal, note = hit
    # 按编号条目拆成多张卡；拆不出则一张综合卡
    parts = re.split(r"\n(?=\d+[）\)])", result)
    items: list[dict[str, Any]] = []
    if len(parts) <= 1:
        item = _optimize_item(
            title="视联网与政务外网运维保障",
            result=result,
            next_goal=next_goal or "继续相关平台设备运维及巡检",
            issue=note or "",
        )
        if item:
            items.append(item)
    else:
        for p in parts:
            p = p.strip()
            if not p:
                continue
            # 标题取首句摘要
            first = re.split(r"[；;，,\n]", p)[0]
            title = re.sub(r"^\d+[）\)]\s*", "", first)[:40] or "运维事项"
            item = _optimize_item(
                title=title[:120],
                result=p,
                next_goal=next_goal or "继续相关平台设备运维及巡检",
                issue=note or "",
            )
            if item:
                items.append(item)
    for i, it in enumerate(items):
        it["sort_order"] = i
    return items, "金滨", 0


def parse_wangjie(rows: list[list[Any]]) -> tuple[list[dict[str, Any]], str | None, int]:
    """王捷驻点表：项目 / 详细 / 状态 / 详情。"""
    buckets: dict[str, list[str]] = {}
    last_project = "日常工作"
    last_detail = ""
    for i, row in enumerate(rows):
        if i == 0:
            continue
        cells = [_cell(c) for c in row]
        if len(cells) < 2:
            continue
        if cells[0] == "项目" and "详细" in (cells[1] if len(cells) > 1 else ""):
            continue
        proj = cells[0] or last_project
        detail = cells[1] if len(cells) > 1 else ""
        status = cells[2] if len(cells) > 2 else ""
        info = cells[3] if len(cells) > 3 else ""
        if cells[0]:
            last_project = cells[0]
        if detail:
            last_detail = detail
        title = last_project
        piece = " / ".join(x for x in [last_detail or detail, status, info] if x)
        if not piece:
            continue
        buckets.setdefault(title, []).append(piece)

    items: list[dict[str, Any]] = []
    skipped = 0
    for title, parts in buckets.items():
        result = "；".join(parts)[:4000]
        issue = ""
        if "离线" in result:
            issue = "存在设备离线，需跟进恢复"
        item = _optimize_item(
            title=title[:120],
            result=result,
            next_goal="持续保障驻点运维与告警处置",
            issue=issue,
        )
        if item:
            items.append(item)
        else:
            skipped += 1
    for i, it in enumerate(items):
        it["sort_order"] = i
    return items, "王捷", skipped


def parse_file(path: Path) -> dict[str, Any]:
    hint = _guess_author_from_filename(path.name)

    # 金滨特殊周志 xls
    if hint == "金滨" and path.suffix.lower() == ".xls":
        items, author, skipped = parse_jinbin(path)
        return {
            "file": path.name,
            "sheet": "2026金滨",
            "author": author or hint,
            "items": items,
            "skipped": skipped,
        }

    if path.suffix.lower() == ".xls":
        rows, sheet = _read_xls_rows(path)
    else:
        rows, sheet = _read_xlsx_rows(path)

    # 王捷特殊表
    if hint == "王捷" or (rows and "驻点" in _cell(rows[0][0] if rows[0] else "")):
        items, author, skipped = parse_wangjie(rows)
        return {
            "file": path.name,
            "sheet": sheet,
            "author": author or hint,
            "items": items,
            "skipped": skipped,
        }

    cols = _detect_header(rows)
    if not cols:
        return {
            "file": path.name,
            "sheet": sheet,
            "author": hint,
            "items": [],
            "skipped": 0,
            "error": "未识别表头",
        }

    items, author, skipped = parse_standard_rows(rows, cols)
    return {
        "file": path.name,
        "sheet": sheet,
        "author": author or hint,
        "items": items,
        "skipped": skipped,
    }


def _fix_items(
    items: list[dict[str, Any]], *, project_id_for_validate: int
) -> list[dict[str, Any]]:
    """校验不过时自动微调，保证可提交。历史导入先挂项目 id 以满足必选校验。"""
    seeded = []
    for it in items:
        cur = dict(it)
        cur["project_id"] = project_id_for_validate
        seeded.append(cur)
    items = seeded
    ok, hard, _soft = validate_report_items(items)
    if ok:
        return items
    fixed: list[dict[str, Any]] = []
    for it in items:
        cur = dict(it)
        cur["project_id"] = project_id_for_validate
        if cur.get("no_action"):
            if not (cur.get("task_title") or "").strip():
                cur["task_title"] = "未命名任务"
            fixed.append(cur)
            continue
        if _is_vague(cur.get("result_text") or ""):
            cur["no_action"] = True
            cur["result_text"] = "本周无实质动作（导入时判定为空话）"
            cur["next_goal"] = ""
            cur["next_due_date"] = None
            cur["stuck_status"] = "none"
            fixed.append(cur)
            continue
        if _is_vague(cur.get("next_goal") or "") or not (cur.get("next_goal") or "").strip():
            cur["next_goal"] = f"继续推进「{cur.get('task_title') or '工作'}」"
        if not cur.get("next_due_date"):
            cur["next_due_date"] = DEFAULT_NEXT_DUE
        if cur.get("stuck_status") in ("blocker", "coordination"):
            if not (cur.get("stuck_detail") or "").strip():
                cur["stuck_detail"] = "需协调（原周报未写清）"
            if not (cur.get("stuck_next_action") or "").strip():
                cur["stuck_next_action"] = cur.get("next_goal") or "跟进协调"
            if not cur.get("stuck_follow_up"):
                cur["stuck_follow_up"] = DEFAULT_NEXT_DUE
        fixed.append(cur)

    # 若全是 no_action，保留至少一条并强制有标题
    if fixed and all(x.get("no_action") for x in fixed):
        fixed[0]["no_action"] = False
        if _is_vague(fixed[0].get("result_text") or ""):
            fixed[0]["result_text"] = "本周按原周报推进相关工作（导入补全）"
        fixed[0]["next_goal"] = fixed[0].get("next_goal") or "按计划继续推进"
        fixed[0]["next_due_date"] = DEFAULT_NEXT_DUE

    ok2, hard2, _ = validate_report_items(fixed)
    if not ok2:
        raise RuntimeError(f"校验仍失败: {hard2}")
    return fixed


async def import_all() -> None:
    if not REPORT_DIR.is_dir():
        raise SystemExit(f"目录不存在: {REPORT_DIR}")

    files = sorted(
        [p for p in REPORT_DIR.iterdir() if p.suffix.lower() in (".xlsx", ".xls")],
        key=lambda p: p.name,
    )
    print(f"发现 {len(files)} 个文件 → {REPORT_DIR}")

    async with async_session() as db:
        daily = await ensure_daily_project(db)
        users = (
            await db.execute(select(User).where(User.is_active.is_(True)))
        ).scalars().all()
        by_name = {u.name: u for u in users}

        summary: list[dict[str, Any]] = []
        for path in files:
            parsed = parse_file(path)
            author_name = parsed.get("author")
            if not author_name or author_name not in by_name:
                summary.append(
                    {
                        "file": path.name,
                        "ok": False,
                        "error": f"未匹配用户: {author_name!r}",
                    }
                )
                print(f"FAIL {path.name}: 未匹配用户 {author_name!r}")
                continue
            user = by_name[author_name]
            try:
                items = _fix_items(
                    list(parsed["items"]), project_id_for_validate=daily.id
                )
            except Exception as exc:
                summary.append({"file": path.name, "ok": False, "error": str(exc)})
                print(f"FAIL {path.name}: {exc}")
                continue

            # 幂等：同用户同周期替换
            existing = await db.execute(
                select(ProgressReport)
                .options(selectinload(ProgressReport.items))
                .where(
                    ProgressReport.user_id == user.id,
                    ProgressReport.period_type == "week",
                    ProgressReport.period_start == PERIOD_START,
                )
            )
            report = existing.scalar_one_or_none()
            if report:
                for old in list(report.items):
                    await db.delete(old)
                await db.flush()
            else:
                report = ProgressReport(
                    user_id=user.id,
                    period_type="week",
                    period_start=PERIOD_START,
                    period_end=PERIOD_END,
                )
                db.add(report)
                await db.flush()

            theme = compose_theme_from_titles(items)
            report.period_end = PERIOD_END
            report.theme = theme
            report.status = "submitted"
            report.submitted_at = datetime.now(timezone.utc)

            for idx, raw in enumerate(items):
                db.add(
                    ProgressReportItem(
                        report_id=report.id,
                        project_id=daily.id,
                        chain_id=None,
                        linked_task_id=None,
                        task_title=(raw.get("task_title") or "")[:120],
                        no_action=bool(raw.get("no_action")),
                        result_text=raw.get("result_text") or "",
                        quant_value=raw.get("quant_value") or "",
                        stuck_status=raw.get("stuck_status") or "none",
                        stuck_detail=raw.get("stuck_detail") or "",
                        stuck_next_action=raw.get("stuck_next_action") or "",
                        stuck_follow_up=raw.get("stuck_follow_up"),
                        stuck_need_who=raw.get("stuck_need_who") or "",
                        linked_blocker_id=None,
                        next_goal=raw.get("next_goal") or "",
                        next_due_date=raw.get("next_due_date"),
                        risk_level=raw.get("risk_level") or "normal",
                        risk_reason=raw.get("risk_reason") or "",
                        risk_impact_date=raw.get("risk_impact_date"),
                        risk_coordinator=raw.get("risk_coordinator") or "",
                        prior_done=None,
                        sort_order=raw.get("sort_order") or idx,
                    )
                )

            active = sum(1 for x in items if not x.get("no_action"))
            summary.append(
                {
                    "file": path.name,
                    "ok": True,
                    "user": user.username,
                    "name": user.name,
                    "dingtalk": bool(user.dingtalk_user_id),
                    "items": len(items),
                    "active": active,
                    "skipped_rows": parsed.get("skipped", 0),
                    "theme": theme[:60],
                }
            )
            print(
                f"OK {user.name}({user.username}) items={len(items)} active={active} "
                f"skip_rows={parsed.get('skipped', 0)} ding={bool(user.dingtalk_user_id)}"
            )

        await db.flush()
        bf = await backfill_null_report_projects(db)
        print(f"backfill_projects={bf}")
        await db.commit()

    print("\n=== SUMMARY ===")
    ok_n = sum(1 for s in summary if s.get("ok"))
    print(f"success={ok_n}/{len(summary)}")
    for s in summary:
        if not s.get("ok"):
            print("  ERR", s)


if __name__ == "__main__":
    import asyncio

    asyncio.run(import_all())
